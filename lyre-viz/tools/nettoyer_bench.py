# Liste (et, sur demande, supprime) les resultats de banc devenus inutiles dans bench\.
#
#   python tools\nettoyer_bench.py                 liste seulement (rien n'est supprime)
#   python tools\nettoyer_bench.py --supprimer     supprime definitivement ce qui est liste
#   python tools\nettoyer_bench.py --jours 7       garde aussi tout ce qui a moins de 7 jours (2 par defaut)
#
# Sont toujours gardes :
# - les dossiers cites dans PERFORMANCE.md, README.md, PLAN-LASERS.md, CLAUDE.md (bench\nom...) ;
# - les dossiers des mesures de l'historique (colonne "Dossiers" de historique_perf.xlsx) ;
# - reference*, lasers_images, et les executables de reference avant*.exe de bench\ref ;
# - tout ce qui a ete modifie depuis moins de --jours jours (travail en cours).
import argparse
import os
import re
import shutil
import sys
import time

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BENCH = os.path.join(ROOT, 'bench')
DOCS = ['PERFORMANCE.md', 'README.md', 'PLAN-LASERS.md', 'CLAUDE.md']
TOUJOURS = {'ref', 'lasers_images'}


def noms_cites():
    noms = set()
    for f in DOCS:
        p = os.path.join(ROOT, f)
        if os.path.exists(p):
            texte = open(p, encoding='utf-8', errors='replace').read()
            for m in re.finditer(r'bench[\\/]+([A-Za-z0-9_\-]+)(\*)?', texte):
                noms.add(m.group(1) + ('*' if m.group(2) else ''))
    try:
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(ROOT, 'historique_perf.xlsx'), read_only=True)
        for ligne in wb['Mesures'].iter_rows(min_row=2, values_only=True):
            if len(ligne) > 14 and ligne[14]:
                noms.update(d.strip() for d in str(ligne[14]).split(','))
    except Exception as e:  # historique absent ou ouvert : on garde tout ce qui pourrait y etre
        print(f'Historique illisible ({e}) : rien ne sera propose a la suppression.')
        sys.exit(1)
    return noms


def garde(nom, cites):
    if nom in TOUJOURS or nom.startswith('reference') or nom in cites:
        return True
    # "bench\essai_*" dans un document : garde tout ce qui commence par "essai_"
    return any(c.endswith('*') and nom.startswith(c[:-1]) for c in cites)


def taille(p):
    if os.path.isfile(p):
        return os.path.getsize(p)
    t = 0
    for d, _, fichiers in os.walk(p):
        for f in fichiers:
            try:
                t += os.path.getsize(os.path.join(d, f))
            except OSError:
                pass
    return t


def recent(p, limite):
    if os.path.getmtime(p) >= limite:
        return True
    for d, _, fichiers in os.walk(p):
        for f in fichiers:
            if os.path.getmtime(os.path.join(d, f)) >= limite:
                return True
    return False


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--supprimer', action='store_true', help='supprime definitivement les elements listes')
    ap.add_argument('--jours', type=float, default=2.0, help='garde ce qui a ete modifie depuis moins de N jours')
    a = ap.parse_args()
    cites = noms_cites()
    limite = time.time() - a.jours * 86400
    candidats = []
    for nom in sorted(os.listdir(BENCH)):
        p = os.path.join(BENCH, nom)
        if garde(nom, cites) or recent(p, limite):
            continue
        candidats.append(p)
    ref = os.path.join(BENCH, 'ref')
    if os.path.isdir(ref):
        for nom in sorted(os.listdir(ref)):
            p = os.path.join(ref, nom)
            if nom.startswith('avant') or os.path.getmtime(p) >= limite:
                continue
            candidats.append(p)
    total = 0
    for p in candidats:
        t = taille(p)
        total += t
        print(f'{t / 1e6:9.0f} Mo  {os.path.relpath(p, ROOT)}')
    print(f'\n{len(candidats)} elements, {total / 1e9:.1f} Go')
    if not a.supprimer:
        print('Rien n\'a ete supprime. Pour supprimer : python tools\\nettoyer_bench.py --supprimer')
        return
    for p in candidats:
        if os.path.isdir(p):
            shutil.rmtree(p)
        else:
            os.remove(p)
    print('Supprime.')


if __name__ == '__main__':
    main()
