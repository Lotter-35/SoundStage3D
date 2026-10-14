# Ajoute les resultats d'un banc d'essai a l'historique des performances (historique_perf.xlsx).
#
#   python tools\historique.py --modif "Ombres des obstacles" --reglages "adaptive=0, en mouvement, shadows=1" bench\essai_B0 bench\essai_B1
#
# - Plusieurs dossiers = plusieurs lancements de la MEME version et des MEMES reglages : leurs valeurs
#   sont moyennees (lancements en alternance, voir PERFORMANCE.md).
# - Une ligne par cas du banc dans la feuille "Mesures" (jamais modifiee ensuite : on ajoute).
# - La feuille "Comparaison" est refaite a chaque ajout : un cas par ligne, une modification par colonne
#   (formules qui lisent "Mesures"), pour voir l'evolution des img/s et du temps GPU.
# - Version : commit git courant, "+modifs" si lyre-viz contient des changements non commites.
# - Banc des lasers (--bench-lasers) : meme format ; colonne "Lasers (ms)" (passes laser) et colonne "Banc"
#   (lyres / lasers) pour distinguer les cas.
import argparse
import datetime
import os
import re
import subprocess
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FICHIER = os.path.join(ROOT, 'historique_perf.xlsx')
COLONNES = ['Date', 'Version', 'Modification', 'Reglages', 'Lancements', 'Cas', 'Image (ms)', 'img/s reelles',
            'GPU moyen (ms)', 'GPU mediane (ms)', 'Surfaces (ms)', 'Fumee (ms)', 'Reflet sol (ms)', "Paires d'ombre",
            'Dossiers', 'Lasers (ms)', 'Banc']
POLICE = 'Arial'


def lire(dossier):
    """Cas -> mesures d'un fichier resultats.txt."""
    cas = {}
    cur = None
    banc = 'lyres'
    with open(os.path.join(dossier, 'resultats.txt'), encoding='utf-8', errors='replace') as f:
        for ligne in f:
            if "banc d'essai lasers" in ligne:
                banc = 'lasers'
            m = re.match(r'\[(\d+)\] (.+)', ligne)
            if m:
                cur = {'nom': m.group(2).strip(), 'surf': 0.0, 'fumee': 0.0, 'reflet': 0.0, 'lasers': 0.0, 'paires': None,
                       'banc': banc}
                cas[int(m.group(1))] = cur
                continue
            if cur is None:
                continue
            m = re.search(r'image : ([\d.]+) ms .*GPU : ([\d.]+) ms \(mediane ([\d.]+)\)', ligne)
            if m:
                cur['image'], cur['gpu'], cur['gpum'] = float(m.group(1)), float(m.group(2)), float(m.group(3))
                continue
            m = re.match(r'\s+(.+?)\s+([\d.]+) ms\s+\(moyenne', ligne)
            if m:
                nom, v = m.group(1).strip(), float(m.group(2))
                if nom.startswith('scene (surfaces)') or nom.startswith('lyres douces'):
                    cur['surf'] += v
                elif nom.startswith('fumee'):
                    cur['fumee'] += v
                elif nom.startswith('reflet'):
                    cur['reflet'] += v
                elif nom.startswith('lasers'):
                    cur['lasers'] += v
                continue
            m = re.search(r'ombres : \d+ lyres, (\d+) paires', ligne)
            if m:
                cur['paires'] = int(m.group(1))
    return cas


def version():
    try:
        h = subprocess.run(['git', 'rev-parse', '--short', 'HEAD'], cwd=ROOT, capture_output=True, text=True).stdout.strip()
        sale = subprocess.run(['git', 'status', '--porcelain', '--', '.'], cwd=ROOT, capture_output=True, text=True).stdout.strip()
        return h + (' +modifs' if sale else '')
    except OSError:
        return '?'


def entete(ws, titres, ligne=1):
    for c, t in enumerate(titres, 1):
        cell = ws.cell(row=ligne, column=c, value=t)
        cell.font = Font(name=POLICE, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', start_color='305496')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def ouvrir():
    if os.path.exists(FICHIER):
        wb = load_workbook(FICHIER)
        ws = wb['Mesures']
        # fichiers d'avant les lasers : colonnes ajoutees a la fin (les anciennes lignes restent vides)
        if ws.max_column < len(COLONNES):
            entete(ws, COLONNES)
            ws.column_dimensions[get_column_letter(len(COLONNES) - 1)].width = 11
            ws.column_dimensions[get_column_letter(len(COLONNES))].width = 9
        return wb
    wb = Workbook()
    ws = wb.active
    ws.title = 'Mesures'
    entete(ws, COLONNES)
    ws.freeze_panes = 'A2'
    largeurs = [17, 15, 38, 34, 11, 40, 11, 11, 13, 13, 12, 11, 12, 11, 40, 11, 9]
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    lis = wb.create_sheet('Lisez-moi')
    texte = [
        "Historique des performances de LyreViz (GTX 1080, 3840 x 2160, 200 lyres)",
        "",
        "Une ligne de 'Mesures' par cas du banc d'essai et par modification. On ajoute, on ne modifie pas.",
        "Ajout : python tools\\historique.py --modif \"nom\" --reglages \"...\" bench\\dossier1 bench\\dossier2",
        "  plusieurs dossiers = lancements de la meme version et des memes reglages, moyennes.",
        "'Comparaison' : un cas par ligne, une modification par colonne (formules sur 'Mesures').",
        "",
        "Colonnes :",
        "  Image (ms) : duree moyenne d'une image (= img/s reelles, presentation comprise).",
        "  GPU moyen / mediane : temps GPU total d'une image.",
        "  Surfaces : passe des surfaces + lyres douces. Fumee : toutes les passes de fumee. Reflet sol : passe miroir.",
        "  Paires d'ombre : paires (lyre, obstacle) de l'ombre (vide avant les ombres).",
        "  Lasers : passes des lasers (banc --bench-lasers). Banc : lyres ou lasers.",
        "",
        "A comparer seulement entre lignes aux memes reglages (qualite adaptative, mouvement / immobile).",
        "Les temps varient de 10 a 20 % d'un lancement a l'autre : voir PERFORMANCE.md.",
        "Objectif : 8,3 ms par image (120 img/s).",
    ]
    for i, t in enumerate(texte, 1):
        lis.cell(row=i, column=1, value=t).font = Font(name=POLICE, bold=(i == 1), size=12 if i == 1 else 10)
    lis.column_dimensions['A'].width = 110
    return wb


def comparaison(wb):
    ws = wb['Mesures']
    lignes = list(ws.iter_rows(min_row=2, values_only=True))
    modifs, cas = [], []
    for r in lignes:
        if r[2] and r[2] not in modifs:
            modifs.append(r[2])
        if r[5] and r[5] not in cas:
            cas.append(r[5])
    if 'Comparaison' in wb.sheetnames:
        del wb['Comparaison']
    cp = wb.create_sheet('Comparaison', 1)
    blocs = [('img/s reelles', 'H', '0'), ('GPU moyen (ms)', 'I', '0.00')]
    ligne = 1
    for titre, col, fmt in blocs:
        cp.cell(row=ligne, column=1, value=titre).font = Font(name=POLICE, bold=True, size=12)
        ligne += 1
        entete(cp, ['Cas'] + modifs, ligne)
        cp.row_dimensions[ligne].height = 60
        hdr = ligne
        for nom in cas:
            ligne += 1
            cp.cell(row=ligne, column=1, value=nom).font = Font(name=POLICE)
            for j in range(len(modifs)):
                c = get_column_letter(j + 2)
                f = (f'=IFERROR(AVERAGEIFS(Mesures!${col}:${col},Mesures!$F:$F,$A{ligne},'
                     f'Mesures!$C:$C,{c}${hdr}),"")')
                cell = cp.cell(row=ligne, column=j + 2, value=f)
                cell.number_format = fmt
                cell.font = Font(name=POLICE)
        ligne += 2
    cp.column_dimensions['A'].width = 40
    for j in range(len(modifs)):
        cp.column_dimensions[get_column_letter(j + 2)].width = 16
    cp.freeze_panes = 'B3'


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('dossiers', nargs='+')
    ap.add_argument('--modif', required=True, help='modification mesuree (nom court, une colonne de la comparaison)')
    ap.add_argument('--reglages', default='', help='reglages du banc (adaptive, mouvement, --set...)')
    ap.add_argument('--version', default=None, help='version (commit) ; par defaut le commit courant')
    ap.add_argument('--date', default=None, help='date (AAAA-MM-JJ HH:MM) ; par defaut maintenant')
    a = ap.parse_args()

    runs = [lire(d) for d in a.dossiers]
    wb = ouvrir()
    ws = wb['Mesures']
    date = a.date or datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
    ver = a.version or version()
    n = 0
    for idx in sorted(runs[0]):
        vals = [r[idx] for r in runs if idx in r and 'gpu' in r[idx]]
        if not vals:
            continue
        moy = lambda k: round(sum(v[k] for v in vals) / len(vals), 3)
        paires = vals[0]['paires']
        row = ws.max_row + 1
        ws.append([date, ver, a.modif, a.reglages, len(vals), vals[0]['nom'], moy('image'), None, moy('gpu'), moy('gpum'),
                   moy('surf'), moy('fumee'), moy('reflet') or None, paires, ', '.join(os.path.basename(os.path.normpath(d)) for d in a.dossiers),
                   moy('lasers') or None, vals[0]['banc']])
        ws.cell(row=row, column=8, value=f'=IF(G{row}>0,1000/G{row},"")')
        for c in range(1, len(COLONNES) + 1):
            ws.cell(row=row, column=c).font = Font(name=POLICE)
        for c, fmt in [(7, '0.00'), (8, '0'), (9, '0.00'), (10, '0.00'), (11, '0.00'), (12, '0.00'), (13, '0.00'), (16, '0.00')]:
            ws.cell(row=row, column=c).number_format = fmt
        n += 1
    comparaison(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(FICHIER)
    print(f'{n} cas ajoutes ({a.modif}) -> {FICHIER}')


if __name__ == '__main__':
    sys.exit(main())
